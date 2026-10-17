"""Writes two small Excel files that mirror the real layouts, for testing.

    python access/tools/make_samples.py
"""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = pathlib.Path(__file__).resolve().parent.parent / "samples"

MASTER_HEADERS = ["بوليصة الشحن", "رقم المرجع", "وصف محتويات الشحنة", "رقم المرجع الثانوي"]
STATUS_HEADERS = ["رقم التتبع", "مدينة الوجهة", "آخر حدث", "اسم المرسل إليه"]

READY = "جاهزة للالتقاط"           # IN
RECEIVED = "استلام من العميل"      # IN
DELIVERED = "التسليم النهائي"      # OUT
IN_TRANSIT = "قيد التوصيل"         # unknown -> auto-added as OUT

X = "ثانوية منار العلم الأهلية - مسارات - M174562"
Y = "متوسطة حائل - مسارات - M174100"

# Contents end in a grade code. G<n> = 1 item of that grade; G<n>-L<k> = a loss
# row of k items for that grade (it does not add to the quantity).
MASTER_ROWS = [
    ["EBUXSL2079600001", "71204-1", "نقطة توزيع المجاردة \\\\ G1", "ASR-SEP30-G1"],
    ["EBUXSL2079600002", "71204-2", "نقطة توزيع المجاردة \\\\ G1", "ASR-SEP30-G1"],
    ["EBUXSL2079600003", "71204-3", "نقطة توزيع المجاردة \\\\ G1", "ASR-SEP30-G1"],
    ["EBUXSL2079600004", "71204-4", "نقطة توزيع المجاردة \\\\ G1", "ASR-SEP30-G1"],
    ["EBUXSL2079600005", "71204-5", "نقطة توزيع المجاردة \\\\ G1", "ASR-SEP30-G1"],
    ["EBUXSL2079600006", "71204-6", "نقطة توزيع المجاردة \\\\ G1-L4", "ASR-SEP30-G1"],
    ["EBUXSL2079700001", "71205-1", "نقطة توزيع حائل // G9", "ASR-SEP30-G9"],
    ["EBUXSL2079700002", "71205-2", "نقطة توزيع حائل / G9-L1", "ASR-SEP30-G9"],
    ["EBUXSL2079700003", "71205-3", "نقطة توزيع حائل \\\\ G9", "ASR-SEP30-G9"],
]

STATUS_ROWS = [
    ["EBUXSL2079600001", "المجاردة", READY, X],
    ["EBUXSL2079600002", "المجاردة", RECEIVED, X],
    ["EBUXSL2079600003", "المجاردة", DELIVERED, X],
    ["EBUXSL2079600004", "المجاردة", DELIVERED, X],
    ["EBUXSL2079600005", "المجاردة", IN_TRANSIT, X],
    ["EBUXSL2079600006", "المجاردة", DELIVERED, X],
    ["EBUXSL2079700001", "حائل", READY, Y],
    ["EBUXSL2079700002", "حائل", DELIVERED, Y],
    ["EBUXSL2079999999", "الرياض", DELIVERED, "مدرسة غير مسجلة - M999999"],  # not in master
    ["EBUXSL2079600004", "المجاردة", DELIVERED, X],                           # duplicate row
]


def write(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for row in rows:
        ws.append(row)
    for col in "ABCD":
        ws.column_dimensions[col].width = 32
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    write(OUT / "master_inventory_sample.xlsx", MASTER_HEADERS, MASTER_ROWS)
    write(OUT / "status_update_sample.xlsx", STATUS_HEADERS, STATUS_ROWS)
    print("written to", OUT)
